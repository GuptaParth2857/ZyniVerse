#!/usr/bin/env python3
"""In-place rebuild of slides 8-12 of ZyniVerse_AI_Presentation.pptx.
Slides 1-7 are left untouched. Same anime theme as the generator."""
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = r"C:\Users\Parth Gupta\Desktop\ZyniVerse\ZyniVerse_AI_Presentation.pptx"
BACKUP = r"C:\Users\PARTHG~1\AppData\Local\Temp\opencode\ZyniVerse_AI_Presentation.backup.pptx"
shutil.copyfile(SRC, BACKUP)

NAVY = RGBColor(0x0F, 0x0E, 0x17)
NAVY2 = RGBColor(0x13, 0x12, 0x22)
PANEL = RGBColor(0x1C, 0x1A, 0x33)
PANEL2 = RGBColor(0x24, 0x21, 0x41)
GRAD = RGBColor(0x2A, 0x27, 0x56)
VIOLET = RGBColor(0x7F, 0x5A, 0xF0)
PINK = RGBColor(0xE5, 0x31, 0x70)
TEAL = RGBColor(0x00, 0xC2, 0xFF)
GOLD = RGBColor(0xFF, 0xC2, 0x57)
GREEN = RGBColor(0x2C, 0xB6, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFE)
MUTED = RGBColor(0xA7, 0xA9, 0xBE)
DIM = RGBColor(0x5A, 0x5B, 0x7A)
LINE = RGBColor(0x2C, 0x2A, 0x48)

F_TITLE = "Segoe UI Black"
F_BODY = "Segoe UI"
F_MONO = "Consolas"


def set_fill(shape, color, alpha=100):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if alpha < 100:
        sf = shape.fill.fore_color._xFill
        srgb = sf.find(qn("a:srgbClr"))
        a = sf.makeelement(qn("a:alpha"), {"val": str(int(alpha * 1000))})
        srgb.append(a)


def no_line(shape):
    shape.line.fill.background()


def shadowless(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def rect(s, x, y, w, h, fill=PANEL, line=None, round_=False, radius=0.12, alpha=100):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    set_fill(sh, fill, alpha)
    if line is not None:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    else:
        no_line(sh)
    shadowless(sh)
    return sh


def oval(s, x, y, w, h, fill, alpha=100, line=None, lw=1.5):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        set_fill(sh, fill, alpha)
    if line is not None:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        no_line(sh)
    shadowless(sh)
    return sh


def _tracks(run, pts):
    try:
        run.font._rPr.set("spc", str(int(pts * 100)))
    except Exception:
        pass


def text(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = para.get("align", align)
        p.space_before = Pt(para.get("sb", 0))
        p.space_after = Pt(para.get("sa", 0))
        p.line_spacing = para.get("ls", 1.0)
        for r in para["runs"]:
            run = p.add_run()
            run.text = r[0]
            f = run.font
            f.size = Pt(r[1])
            f.color.rgb = r[2]
            f.bold = r[3] if len(r) > 3 else False
            f.name = r[4] if len(r) > 4 else F_BODY
            if len(r) > 5 and r[5]:
                f.italic = True
            if len(r) > 6 and r[6]:
                _tracks(run, r[6])
    return tb


def shape_text(sh, paras, anchor=MSO_ANCHOR.MIDDLE, ml=0.22, mr=0.22):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(ml)
    tf.margin_right = Inches(mr)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = para.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(para.get("sb", 0))
        p.space_after = Pt(para.get("sa", 0))
        p.line_spacing = para.get("ls", 1.0)
        for r in para["runs"]:
            run = p.add_run()
            run.text = r[0]
            f = run.font
            f.size = Pt(r[1])
            f.color.rgb = r[2]
            f.bold = r[3] if len(r) > 3 else False
            f.name = r[4] if len(r) > 4 else F_BODY
            if len(r) > 5 and r[5]:
                f.italic = True
            if len(r) > 6 and r[6]:
                _tracks(run, r[6])
    return sh


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


def decorate(s):
    oval(s, -1.6, -1.9, 5.2, 5.2, VIOLET, alpha=6)
    oval(s, 10.6, 4.9, 4.6, 4.6, PINK, alpha=5)
    oval(s, 11.3, -2.4, 3.4, 3.4, TEAL, alpha=5)
    text(s, 11.95, 0.42, 1.0, 0.5, [{"runs": [("\u2726", 18, PINK, True, F_BODY)]}])
    text(s, 11.6, 0.95, 0.8, 0.4, [{"runs": [("\u2727", 11, GOLD, True, F_BODY)]}])
    oval(s, 12.62, 1.14, 0.09, 0.09, TEAL)
    oval(s, 0.55, 6.92, 0.06, 0.06, VIOLET)


def base_slide(s):
    rect(s, 0, 0, 13.333, 7.5, NAVY)
    decorate(s)


def header(s, idx, title, subtitle=None, accent=PINK):
    chip = rect(s, 0.56, 0.52, 0.52, 0.52, VIOLET, round_=True, radius=0.22)
    shape_text(chip, [{"runs": [(str(idx), 20, WHITE, True, F_TITLE)], "align": PP_ALIGN.CENTER}], ml=0, mr=0)
    text(s, 1.28, 0.42, 11.2, 0.75, [{"runs": [(title, 27, WHITE, True, F_TITLE)]}])
    if subtitle:
        text(s, 1.30, 1.18, 11.2, 0.4, [{"runs": [(subtitle, 12.5, MUTED, False, F_BODY)]}])
    rect(s, 1.30, 1.66, 1.6, 0.05, VIOLET)
    rect(s, 2.95, 1.66, 0.7, 0.05, PINK)


def brand(s, idx):
    rect(s, 0.56, 6.88, 12.21, 0.016, LINE)
    text(s, 0.56, 7.0, 6.0, 0.35, [{"runs": [("ZYNIVERSE \u00d7 AI", 9.5, DIM, True, F_BODY, False, 3.2)]}])
    text(s, 7.5, 7.0, 5.27, 0.35, [{"runs": [("PARTH GUPTA \u00b7 2026 \u00b7 " + str(idx) + " / 12", 9.5, DIM, False, F_BODY, False, 1.6)],
                                    "align": PP_ALIGN.RIGHT}])


def code_strip(s, fname, lines, x=0.56, y=6.16, w=12.21, title_color=TEAL):
    rect(s, x, y, w, 0.62, PANEL, round_=True, radius=0.14)
    rect(s, x, y, 0.05, 0.62, VIOLET)
    text(s, x + 0.22, y + 0.075, 6.0, 0.25, [{"runs": [("CODE \u2014 " + fname, 10, title_color, True, F_MONO)]}])
    row = "   \u00b7   ".join(lines) if lines else "\u2026"
    n = len(row)
    size = 9.0 if n > 165 else (9.5 if n > 140 else (10.0 if n > 115 else 10.5))
    text(s, x + 0.22, y + 0.315, w - 0.5, 0.28, [{"runs": [(row, size, MUTED, False, F_MONO)]}])


prs = Presentation(SRC)
slides = [s for s in prs.slides]


def clear(s):
    for sh in list(s.shapes):
        el = sh._element
        el.getparent().remove(el)


def rebuild(idx, builder):
    sl = slides[idx - 1]
    clear(sl)
    builder(sl)


# ============================================================
# SLIDE 8 — SHARED WATCHING & SOCIAL AI
# ============================================================
def build8(sl):
    base_slide(sl)
    header(sl, 8, "Shared Watching + Social AI", "Planning, co-watching and debate — all inside one platform", PINK)
    cards = [
        ("Night Planner", VIOLET, "\u263e",
         "Blends every friend's taste into a themed lineup of anime, snacks & trivia \u2014 then auto-creates the watch party."),
        ("Watch Parties", TEAL, "\u25b6",
         "Socket.io co-viewing: play / pause / seek sync, WebRTC voice chat + screen-share, HLS playback."),
        ("Character Tournaments", PINK, "\u2605",
         "Crowd-ranked character battles, polls and Zyni Awards \u2014 community votes drive the results."),
        ("AI Lab \u2014 Fan Theories", GOLD, "\u26a1",
         "AI generates theories per anime; the fandom agrees or disagrees \u2014 theories become content."),
    ]
    for i, (t, c, ic, d) in enumerate(cards):
        r, cl = divmod(i, 2)
        x = 0.56 + cl * 6.28
        y = 1.95 + r * 1.95
        rect(sl, x, y, 5.95, 1.8, PANEL, round_=True, radius=0.07)
        rect(sl, x + 0.2, y + 0.2, 0.06, 0.52, c, round_=True, radius=0.5)
        text(sl, x + 0.44, y + 0.18, 5.0, 0.4, [{"runs": [(ic + "  ", 14, c, True, F_BODY), (t, 14, WHITE, True, F_TITLE)]}])
        text(sl, x + 0.44, y + 0.62, 5.25, 1.05, [{"runs": [(d, 11, MUTED, False, F_BODY)], "ls": 1.1}])
    pill = rect(sl, 0.56, 5.72, 12.21, 0.4, PANEL2, round_=True, radius=0.5)
    shape_text(pill, [{"runs": [("EXAMPLE \u2192  \"Sunday Funday\" \u2014 friends' tastes fused \u2192 Jujutsu Kaisen marath + momos + trivia \u2192 watch-party room live", 11.5, WHITE, True, F_BODY)]}])
    code_strip(sl, "watch-party.ts + ai/night-gen", ["Socket.io: play/pause/seek sync", "WebRTC voice + screen-share", "night-gen() → createParty()", "tournament polls → ranked results"])
    brand(sl, 8)
    notes(sl, "Two AIs + two social engines: Night Planner fuses a friend group's taste into a themed lineup then auto-creates a real watch-party. Watch Parties run on Socket.io with synced playback plus WebRTC voice and screen-share. Character Tournaments give the crowd ranked battles. AI Lab generates fan theories the community then votes on.")


# ============================================================
# SLIDE 9 — PRACTICAL AI + CONTENT TOOLS
# ============================================================
def build9(sl):
    base_slide(sl)
    header(sl, 9, "Practical AI + Content Tools", "Filler that skips itself, reels and manga that build themselves", GREEN)
    cards = [
        ("Skip Filler", TEAL, "\u25b8",
         "AI summarises the filler episode so you skip it guilt-free \u2014 plus crowd-voted filler list and 200 guides."),
        ("AI Search + Mixer", VIOLET, "\u2315",
         "Natural questions + semantic intents layered over the AniList index \u2014 'My Hero Academia but space horror' mixes."),
        ("Hinglish Reels Generator", PINK, "\u270d",
         "ffmpeg + impact font + sharp render Hinglish meme reels automatically from source clips (scripts/make-reels.mjs)."),
        ("In-app Manga Reader", GOLD, "\u25a3",
         "Read chapters right in the browser \u2014 chapter browser, progress tracking, doujinshi & light-novel support."),
    ]
    for i, (t, c, ic, d) in enumerate(cards):
        r, cl = divmod(i, 2)
        x = 0.56 + cl * 6.28
        y = 1.95 + r * 1.95
        rect(sl, x, y, 5.95, 1.8, PANEL, round_=True, radius=0.07)
        rect(sl, x + 0.2, y + 0.2, 0.06, 0.52, c, round_=True, radius=0.5)
        text(sl, x + 0.44, y + 0.18, 5.0, 0.4, [{"runs": [(ic + "  ", 14, c, True, F_BODY), (t, 14, WHITE, True, F_TITLE)]}])
        text(sl, x + 0.44, y + 0.62, 5.25, 1.05, [{"runs": [(d, 11, MUTED, False, F_BODY)], "ls": 1.1}])
    pill = rect(sl, 0.56, 5.72, 12.21, 0.4, PANEL2, round_=True, radius=0.5)
    shape_text(pill, [{"runs": [("PATTERN \u2192  content tooling + AI share one stack: import scripts \u2192 ffmpeg/typed endpoints \u2192 quota \u2192 cache \u2192 typed UI", 11.5, TEAL, True, F_MONO)]}])
    code_strip(sl, "scripts/make-reels.mjs + manga-reader", ["ffmpeg-static + impact font \u2192 Hinglish reels", "filler votes + guides", "chapter browser with progress sync"])
    brand(sl, 9)
    notes(sl, "Practical features, not demo-only: Skip Filler sums up filler episodes (supported by crowd-voted filler data and 200 guides); AI search and the Mixer layer semantics over the real AniList index; the Hinglish reels generator automates content creation with an ffmpeg pipeline; and the in-app manga reader covers chapters, doujinshi and light novels.")


# ============================================================
# SLIDE 10 — RECOMMENDER ENGINE (+ WRAPPED TIE-IN)
# ============================================================
def build10(sl):
    base_slide(sl)
    header(sl, 10, "The AI That Doesn't Need the LLM", "Behavioral recommendation engine \u2014 src/lib/ai-recommender.ts", VIOLET)
    text(sl, 0.56, 1.95, 6.0, 0.4, [{"runs": [("HOW IT SCORES", 13, TEAL, True, F_TITLE)]}])
    steps = [
        ("Profile build", "Last 30 days of UserActivity \u2192 genres, viewed IDs, searches."),
        ("Cold-start seed", "New users get taste seeded from onboarding favoriteGenres."),
        ("Signal weighting", "Actions carry weight: view=3, add-to-list=5, click-rec=4, search=2\u2026"),
        ("Time decay", "Older events shrink \u2014 score \u221d recency (1 \u2212 days/30, floor 0.1)."),
        ("Rank \u2014 recs + Wrapped", "Scored recs with reasons; feeds 'My Year in Anime' recap + streaks."),
    ]
    yy = 2.42
    for t, d in steps:
        rect(sl, 0.56, yy, 6.3, 0.62, PANEL, round_=True, radius=0.14)
        rect(sl, 0.62, yy + 0.09, 0.05, 0.44, TEAL, round_=True, radius=0.5)
        text(sl, 0.92, yy + 0.07, 2.3, 0.45, [{"runs": [(t, 10.5, TEAL, True, F_MONO)], "ls": 1.0}], anchor=MSO_ANCHOR.MIDDLE)
        text(sl, 3.15, yy + 0.06, 3.6, 0.5, [{"runs": [(d, 10, MUTED, False, F_BODY)], "ls": 1.05}], anchor=MSO_ANCHOR.MIDDLE)
        yy += 0.72
    rx = 7.15
    text(sl, rx, 1.95, 5.6, 0.4, [{"runs": [("ACTION WEIGHTS", 13, PINK, True, F_TITLE)]}])
    rows = [("view_anime", "3"), ("add_to_list", "5"), ("click_recommendation", "4"),
            ("search", "2"), ("view_filler", "2"), ("view_character", "1.5")]
    tbh = 0.5
    ty = 2.42
    rect(sl, rx, ty, 5.62, tbh, PANEL2, round_=True, radius=0.2)
    text(sl, rx + 0.3, ty, 3.5, tbh, [{"runs": [("SIGNAL \u2014 USER ACTIVITY", 10.5, WHITE, True, F_MONO, False, 1.5)]}], anchor=MSO_ANCHOR.MIDDLE)
    text(sl, rx + 4.3, ty, 1.2, tbh, [{"runs": [("WT", 10.5, WHITE, True, F_MONO, False, 1.5)], "align": PP_ALIGN.RIGHT}], anchor=MSO_ANCHOR.MIDDLE)
    ty += tbh + 0.04
    for k, v in rows:
        rect(sl, rx, ty, 5.62, 0.46, PANEL, round_=True, radius=0.2)
        text(sl, rx + 0.3, ty, 4.2, 0.46, [{"runs": [(k, 11, MUTED, False, F_MONO)]}], anchor=MSO_ANCHOR.MIDDLE)
        text(sl, rx + 4.3, ty, 1.2, 0.46, [{"runs": [(v, 11.5, GOLD, True, F_MONO)], "align": PP_ALIGN.RIGHT}], anchor=MSO_ANCHOR.MIDDLE)
        ty += 0.5
    code_strip(sl, "ai-recommender.ts", ["Decay = max(0.1, 1 \u2212 days/30)", "genre weights seeded at onboarding", "rank = genre \u00b7 search \u00b7 similarity", "\u2192 recs + My-Year-in-Anime recap"])
    brand(sl, 10)
    notes(sl, "Not all AI here is a language model. This deterministic engine builds a 30-day activity profile, seeds cold-start users from onboarding favorite genres, weights actions (add-to-list=5 vs a plain view=3), applies time-decay so old events fade, and ranks candidates by genre/search/similarity \u2014 each rec ships with a human-readable reason. The same profile powers the 'My Year in Anime' Wrapped recap and streaks.")


# ============================================================
# SLIDE 11 — THE ECOSYSTEM NOBODY ELSE SHIPS
# ============================================================
def build11(sl):
    base_slide(sl)
    header(sl, 11, "The Ecosystem Nobody Else Ships", "Software that lives where the fans already are", PINK)
    cards = [
        ("Chrome Extension (MV3)", VIOLET, "\u2726",
         "Filler % + Indian-dub badge injected right on MAL / AniList / Anime-Planet; right-click \u2192 'Search on ZyniVerse'."),
        ("JioTV EPG TV Schedule", TEAL, "\u25a3",
         "Scraped Indian TV anime schedule (EPG cache) \u2014 know exactly when an episode airs on your channel."),
        ("Discord Bot", PINK, "\u25c9",
         "A fourth surface: notifications + community commands bridge ZyniVerse straight into the fandom."),
        ("Public API v1 + Guides", GOLD, "\u25c6",
         "Developer API \u2014 scrobble, dub-status, filler, schedule, webhooks \u2014 plus 200 filler guides and 40 watch orders."),
    ]
    for i, (t, c, ic, d) in enumerate(cards):
        r, cl = divmod(i, 2)
        x = 0.56 + cl * 6.28
        y = 1.95 + r * 1.95
        rect(sl, x, y, 5.95, 1.8, PANEL, round_=True, radius=0.07)
        rect(sl, x + 0.2, y + 0.2, 0.06, 0.52, c, round_=True, radius=0.5)
        text(sl, x + 0.44, y + 0.18, 5.0, 0.4, [{"runs": [(ic + "  ", 14, c, True, F_BODY), (t, 14, WHITE, True, F_TITLE)]}])
        text(sl, x + 0.44, y + 0.62, 5.25, 1.05, [{"runs": [(d, 11, MUTED, False, F_BODY)], "ls": 1.1}])
    pill = rect(sl, 0.56, 5.72, 12.21, 0.4, PANEL2, round_=True, radius=0.5)
    shape_text(pill, [{"runs": [("STILL HARDENED \u2192 37 scrapers blocked \u00b7 3-layer caching \u00b7 k6 load-tested \u00b7 Stripe Pro + UPI \u00b7 ads/merch/affiliate", 11.5, GOLD, True, F_BODY)]}])
    code_strip(sl, "manifest.json + server.ts + api/v1", ["MV3 content-script \u2192 filler % + dub badge", "JioTV EPG scrape \u2192 tv-schedule cache", "discord-bot notifications", "api/v1: scrobble \u00b7 dub-status \u00b7 schedule"])
    brand(sl, 11)
    notes(sl, "These are the features your class project almost certainly doesn't have: an MV3 Chrome extension that injects filler percentage and Indian dub status directly onto MAL/AniList/Anime-Planet plus right-click search; a JioTV EPG scraper powering the TV schedule; a Discord bot; and a public API v1 with scrobbling, dub status and webhooks. All of it rides production hardening: 37 scrapers blocked, three caching layers, k6 load tests, Stripe Pro with UPI.")


# ============================================================
# SLIDE 12 — CLOSING / DEMO
# ============================================================
def build12(sl):
    base_slide(sl)
    header(sl, 12, "AI Isn't the Feature \u2014 It's the Engine", "One pipeline + one behavioral engine = a personal anime universe", PINK)
    text(sl, 0.56, 1.92, 12.2, 0.4, [{"runs": [("LIVE DEMO FLOW", 13, TEAL, True, F_TITLE, False, 2.0)]}])
    flow = [("Zyni AI Chat", "ask anything", VIOLET),
            ("Screenshot Identify", "drop a frame", TEAL),
            ("Vibe Search", "type a mood", PINK),
            ("Install Extension", "dub badge", GOLD),
            ("Night Planner", "invite friends", GREEN),
            ("Watch Party", "press play", VIOLET)]
    fx = 0.56
    step = 2.07
    cw = 1.85
    for i, (t, d, c) in enumerate(flow):
        rect(sl, fx, 2.42, cw, 1.35, PANEL, round_=True, radius=0.12)
        rect(sl, fx, 2.42, cw, 0.05, c)
        text(sl, fx + 0.12, 2.58, cw - 0.24, 0.6, [{"runs": [(t, 10.5, WHITE, True, F_TITLE)], "ls": 1.0}])
        text(sl, fx + 0.12, 3.16, cw - 0.24, 0.4, [{"runs": [(d, 9, MUTED, False, F_BODY)]}])
        if i < len(flow) - 1:
            text(sl, fx + cw + 0.02, 2.85, 0.2, 0.5, [{"runs": [("\u25b8", 14, GOLD, True, F_BODY)]}])
        fx += step
    rect(sl, 0.56, 4.22, 12.21, 1.0, PANEL2, round_=True, radius=0.1)
    rect(sl, 0.56, 4.22, 0.06, 1.0, VIOLET)
    text(sl, 0.95, 4.36, 11.4, 0.72, [{"runs": [("TAKEAWAY  \u2014  ", 15, PINK, True, F_TITLE),
                                        ("A real community at 224+ pages and ~190 models, where one cost-controlled Gemini pipeline, a behavioral recommender, a Chrome extension and real-time watch parties give every user a personal anime universe.", 14, WHITE, False, F_BODY)], "ls": 1.2}])
    text(sl, 0.56, 5.66, 12.2, 0.75, [{"runs": [("THANK YOU \u00b7 Q & A", 30, WHITE, True, F_TITLE)], "align": PP_ALIGN.CENTER}])
    text(sl, 0.56, 6.42, 12.2, 0.4, [{"runs": [("zyverse.in   \u00b7   github.com <your-username>   \u00b7   Let's argue about best girl after this", 12.5, MUTED, False, F_BODY)], "align": PP_ALIGN.CENTER}])
    brand(sl, 12)
    notes(sl, "Close on the thesis: AI is the engine, not a feature. Walk the live demo: ask Zyni AI, drop a screenshot to identify, search by vibe, install the extension for the dub badge, plan an anime night, then press play together. Thank the judges and open Q&A. Mention zyverse.in.")


for num, builder in [(8, build8), (9, build9), (10, build10), (11, build11), (12, build12)]:
    rebuild(num, builder)

prs.save(SRC)
print("PATCHED:", SRC, "| slides:", len(prs.slides._sldIdLst))